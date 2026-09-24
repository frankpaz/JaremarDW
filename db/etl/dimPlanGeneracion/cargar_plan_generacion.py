"""
Carga masiva del plan de generacion de energia desde el Excel del equipo de
planta ('Plan Produccion Energetica - Calculos.xlsx', hoja Sheet1) hacia
stg -> [int] -> dw.dimPlanGeneracion.

La tabla es FIEL al Excel: solo los 9 campos marcados en verde, con sus nombres
(Planta, Proveedor, Inversor, InversorId, Mes, NombreMes, Ubicacion,
PlanDiarioAsignado, PlanDiarioInversor), mas la fecha del dia y las columnas de
auditoria. Se cargan TODAS las filas del archivo (tambien las de plan 0, ej. SOLIS);
una celda con error de Excel (#DIV/0!) se guarda como NULL, no como 0.
PlanDiarioAsignado es el plan de la PLANTA y se repite en cada inversor de la planta.

El archivo trae el mes, no el anio ni el dia: se indica --anio y cada mes se expande a
sus DIAS REALES (una fila por fecha e inversor). Si el archivo trae ademas
PlanTotalPlanta y PlanDiarioPlanta, comprueba que el divisor sea el numero de dias del
mes (avisa si se uso un divisor fijo, ej. /31).

Valida antes de escribir (cabeceras, Mes 1-12, valores numericos >= 0, sin
inversor-mes duplicado); si hay errores no se escribe nada. Avisa, sin bloquear:
  * inversores con plan 0 aunque su planta tiene plan (capacidad faltante en el archivo);
  * plantas cuyos inversores suman menos que el plan de la planta;
  * inversores que no existen en la dimension de su fabricante (solo revision: no se guardan llaves).

Cada carga entra con una --version y queda en el historial: para cada fecha y dispositivo
queda vigente la version cargada mas recientemente; las fechas que la nueva carga no trae
conservan su version anterior. Recargar la misma version es idempotente.
--reemplaza-version deja como historicas, por completo, las versiones indicadas.
--limpiar vacia dimPlanGeneracion (stg, int y dw, con su historial) justo antes de cargar, solo
despues de validar el archivo; es destructivo y sin el flag nunca se borra nada.

Requiere openpyxl (solo para este cargador).

Uso:
    python db/etl/dimPlanGeneracion/cargar_plan_generacion.py --archivo "<ruta>.xlsx" --anio 2026 --version V1 --dry-run
    python db/etl/dimPlanGeneracion/cargar_plan_generacion.py --archivo "<ruta>.xlsx" --anio 2026 --version V1 [--limpiar] [--reemplaza-version V0] [--env-file .env.prod]
"""
import argparse
import calendar
import datetime
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent.parent.parent
sys.path.insert(0, str(ROOT / "db"))
from migrate import build_connection, load_env  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HOJA_DEFECTO = "Sheet1"
FUENTE_DEFECTO = "Plan Produccion Energetica - Calculos"
REQUERIDAS = ["PLANTA", "PROVEEDOR", "INVERSOR", "INVERSORID", "MES", "NOMBREMES", "UBICACION",
              "PLANDIARIOASIGNADO", "PLANDIARIOINVERSOR"]
TOLERANCIA_REPARTO = 0.005  # 0.5 %


def norm(texto) -> str:
    s = unicodedata.normalize("NFD", str(texto or ""))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().upper()


def es_error_excel(v) -> bool:
    return isinstance(v, str) and v.strip().startswith("#")


def es_numero(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# --- Lectura y validacion del Excel ------------------------------------------

def leer_filas(wb, hoja, errores: list, avisos: list) -> list:
    if hoja not in wb.sheetnames:
        errores.append(f"No existe la hoja '{hoja}'. Hojas: {wb.sheetnames}")
        return []
    filas = list(wb[hoja].iter_rows(values_only=True))
    if not filas:
        errores.append(f"La hoja '{hoja}' esta vacia.")
        return []
    cab = {norm(c): j for j, c in enumerate(filas[0]) if c is not None}
    faltan = [c for c in REQUERIDAS if c not in cab]
    if faltan:
        errores.append(f"Faltan las columnas {faltan} en la primera fila de '{hoja}'.")
        return []
    g = lambda r, k: r[cab[k]] if cab[k] < len(r) else None
    tiene_diario = "PLANTOTALPLANTA" in cab and "PLANDIARIOPLANTA" in cab

    resultado, vistos, celdas_error = [], set(), 0
    for n, r in enumerate(filas[1:], start=2):
        if r is None or all(v is None for v in r):
            continue
        texto = {k: str(g(r, k)).strip() if g(r, k) is not None else "" for k in ("PLANTA", "PROVEEDOR", "INVERSOR", "INVERSORID", "NOMBREMES", "UBICACION")}
        vacios = [k for k, v in texto.items() if not v]
        if vacios:
            errores.append(f"'{hoja}' fila {n}: faltan valores en {vacios}.")
            continue
        mes = g(r, "MES")
        if not es_numero(mes) or int(mes) != mes or not 1 <= int(mes) <= 12:
            errores.append(f"'{hoja}' fila {n} ({texto['INVERSORID']}): Mes invalido {mes!r}.")
            continue
        mes = int(mes)
        if (texto["INVERSORID"], mes) in vistos:
            errores.append(f"'{hoja}' fila {n}: el inversor {texto['INVERSORID']} tiene el mes {mes} duplicado.")
            continue
        vistos.add((texto["INVERSORID"], mes))
        valores = {}
        for clave, nombre in (("PLANDIARIOINVERSOR", "PlanDiarioInversor"), ("PLANDIARIOASIGNADO", "PlanDiarioAsignado")):
            v = g(r, clave)
            if es_error_excel(v) or v is None:
                celdas_error += 1
                valores[nombre] = None
            elif not es_numero(v):
                errores.append(f"'{hoja}' fila {n} ({texto['INVERSORID']}), {nombre}: valor no numerico {v!r}.")
                valores[nombre] = None
            elif v < 0:
                errores.append(f"'{hoja}' fila {n} ({texto['INVERSORID']}), {nombre}: valor negativo {v}.")
                valores[nombre] = None
            else:
                valores[nombre] = round(float(v), 8)
        dias_impl = None
        if tiene_diario and es_numero(g(r, "PLANTOTALPLANTA")) and es_numero(g(r, "PLANDIARIOPLANTA")) and g(r, "PLANDIARIOPLANTA"):
            dias_impl = g(r, "PLANTOTALPLANTA") / g(r, "PLANDIARIOPLANTA")
        resultado.append({
            "Planta": texto["PLANTA"], "Proveedor": texto["PROVEEDOR"].lower(), "Inversor": texto["INVERSOR"],
            "InversorId": texto["INVERSORID"], "Mes": mes, "NombreMes": texto["NOMBREMES"], "Ubicacion": texto["UBICACION"],
            "PlanDiarioAsignado": valores["PlanDiarioAsignado"], "PlanDiarioInversor": valores["PlanDiarioInversor"],
            "DiasImplicitos": dias_impl,
        })
    if not resultado:
        errores.append(f"'{hoja}' no tiene filas de datos.")
    if celdas_error:
        avisos.append(f"{celdas_error} celda(s) de plan vacias o con error de Excel; se guardan como NULL (no como 0).")
    return resultado


def evaluar_calidad(filas: list, anio: int, avisos: list) -> None:
    """Agrega los avisos de calidad; no filtra ni modifica filas."""
    malos = sorted({f["Mes"] for f in filas if f["DiasImplicitos"] is not None
                    and abs(f["DiasImplicitos"] - calendar.monthrange(anio, f["Mes"])[1]) > 0.01})
    if malos:
        avisos.append(f"PlanDiarioPlanta no corresponde a los dias reales de {anio} en los meses {malos} "
                      "(PlanTotalPlanta / PlanDiarioPlanta no es el numero de dias del mes; posible divisor fijo).")

    sin_plan, cero_inv = {}, {}
    for f in filas:
        pi, pa = f["PlanDiarioInversor"], f["PlanDiarioAsignado"]
        if (pi in (0, None)) and (pa in (0, None)):
            sin_plan.setdefault(f["Planta"], set()).add(f["InversorId"])
        elif pi == 0 and pa and pa > 0:
            cero_inv.setdefault(f["Planta"], set()).add(f["InversorId"])
    for planta, invs in sorted(sin_plan.items()):
        avisos.append(f"{planta}: plan 0 o vacio en todos sus inversores y en la planta ({sorted(invs)}); se carga tal cual.")
    for planta, invs in sorted(cero_inv.items()):
        avisos.append(f"{planta}: {len(invs)} inversor(es) con plan 0 aunque la planta tiene plan asignado "
                      f"({sorted(invs)}); suele ser capacidad faltante en el archivo.")

    por = {}
    for f in filas:
        if f["PlanDiarioInversor"] is None or f["PlanDiarioAsignado"] is None:
            continue
        d = por.setdefault((f["Planta"], f["Mes"]), {"asig": f["PlanDiarioAsignado"], "suma": 0.0})
        d["suma"] += f["PlanDiarioInversor"]
    resumen = {}
    for (planta, mes), d in por.items():
        if d["asig"]:
            r = resumen.setdefault(planta, [0.0, 0.0])
            r[0] += d["suma"]
            r[1] += d["asig"]
    for planta, (suma, asig) in sorted(resumen.items()):
        if asig and abs(suma / asig - 1) > TOLERANCIA_REPARTO:
            avisos.append(f"{planta}: los inversores suman {suma / asig:.1%} del plan diario de la planta "
                          "(el resto no esta repartido entre inversores).")


def expandir_dias(filas: list, anio: int) -> list:
    salida = []
    for f in filas:
        for dia in range(1, calendar.monthrange(anio, f["Mes"])[1] + 1):
            salida.append({**f, "Fecha": datetime.date(anio, f["Mes"], dia)})
    return salida


def avisar_dispositivos(cur, filas: list, avisos: list) -> None:
    """Revision (no se guardan llaves): inversores que no existen en la dimension de su fabricante."""
    consultas = {
        "sma": "SELECT CAST(DeviceId AS NVARCHAR(50)) FROM dw.dimSmaDevices",
        "huawei": "SELECT CAST(DeviceId AS NVARCHAR(50)) FROM dw.dimHuaweiDevices",
        "soliscloud": "SELECT DeviceSn FROM dw.dimSoliscloudDevices",
        "growatt": "SELECT JoinKey FROM dw.dimDeviceCapacity WHERE JoinKey IS NOT NULL",
    }
    try:
        conocidos = {}
        for prov, sql in consultas.items():
            cur.execute(sql)
            conocidos[prov] = {str(r[0]).strip() for r in cur.fetchall() if r[0] is not None}
    except Exception:
        avisos.append("No se pudieron leer las dimensiones de dispositivos; se omite la revision.")
        return
    sin = sorted({(f["Proveedor"], f["InversorId"]) for f in filas if f["Inversor"] not in conocidos.get(f["Proveedor"], set())})
    if sin:
        avisos.append(f"{len(sin)} inversor(es) que no existen en la dimension de su fabricante: {sin}.")


# --- Carga (framework de control ETL) -----------------------------------------

def ejecutar_paso(conn, proceso, origen, destino, fn):
    cur = conn.cursor()
    cur.execute(
        """
        DECLARE @pid INT;
        EXEC dbo.usp_Etl_ProcesoRegistrar
            @Proceso = ?, @Dominio = ?, @SistemaOrigen = ?,
            @EsquemaOrigen = ?, @TablaOrigen = ?,
            @EsquemaDestino = ?, @TablaDestino = ?, @TipoCarga = ?,
            @ProcesoId = @pid OUTPUT;
        SELECT @pid;
        """,
        proceso, "PlanGeneracion", origen[0], origen[0], origen[1], destino[0], destino[1], "FULL",
    )
    cur.fetchone()
    conn.commit()
    cur.execute("EXEC dbo.usp_Etl_RunIniciar @Proceso = ?", proceso)
    run_id = cur.fetchone()[0]
    conn.commit()
    try:
        leidas, insertadas, actualizadas, ignoradas = fn(cur, run_id)
        print(f"  [{proceso}] leidas {leidas} / insertadas {insertadas} / actualizadas {actualizadas} / sin cambio {ignoradas}")
        cur.execute(
            "EXEC dbo.usp_Etl_RunFinalizar @RunId = ?, @Estado = ?, @FilasLeidas = ?, @FilasInsertadas = ?, "
            "@FilasActualizadas = ?, @FilasIgnoradas = ?, @MensajeError = ?, @TareaError = ?",
            run_id, "EXITO", leidas, insertadas, actualizadas, ignoradas, None, None,
        )
        cur.execute(
            "EXEC dbo.usp_Etl_WatermarkActualizar @Proceso = ?, @NuevaFechaHora = ?, @TipoCarga = ?",
            proceso, datetime.datetime.now(), "FULL",
        )
        conn.commit()
    except Exception as exc:
        conn.rollback()
        cur.execute(
            "EXEC dbo.usp_Etl_RunFinalizar @RunId = ?, @Estado = ?, @FilasLeidas = ?, @FilasInsertadas = ?, "
            "@FilasActualizadas = ?, @FilasIgnoradas = ?, @MensajeError = ?, @TareaError = ?",
            run_id, "ERROR", None, None, None, None, str(exc), proceso,
        )
        conn.commit()
        raise


def cargar_stg(plan, archivo, fuente, version):
    def fn(cur, run_id):
        cur.execute("TRUNCATE TABLE stg.dimPlanGeneracion")
        cur.fast_executemany = True
        cur.executemany(
            "INSERT INTO stg.dimPlanGeneracion (Fecha, Planta, Proveedor, Inversor, InversorId, Mes, NombreMes, Ubicacion, "
            "PlanDiarioAsignado, PlanDiarioInversor, PlanFuente, PlanVersion, ArchivoOrigen, RunId) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [(p["Fecha"], p["Planta"], p["Proveedor"], p["Inversor"], p["InversorId"], p["Mes"], p["NombreMes"], p["Ubicacion"],
              p["PlanDiarioAsignado"], p["PlanDiarioInversor"], fuente, version, archivo, run_id) for p in plan],
        )
        return len(plan), len(plan), 0, 0
    return fn


def limpiar_tablas(cur, run_id):
    """Vacia dw, int y stg de dimPlanGeneracion. Devuelve (filas dw + int borradas, 0, 0, 0)."""
    n = 0
    for tabla in ("dw.dimPlanGeneracion", "[int].dimPlanGeneracion", "stg.dimPlanGeneracion"):
        cur.execute(f"SELECT COUNT(*) FROM {tabla}")
        filas = cur.fetchone()[0]
        if not tabla.startswith("stg"):
            n += filas
        cur.execute(f"TRUNCATE TABLE {tabla}")
        print(f"    {tabla}: {filas} filas eliminadas")
    return n, 0, 0, 0


def retirar_versiones(versiones, version_nueva):
    """Marca como historicas TODAS las filas de las versiones indicadas (silver; gold lo hereda en su merge)."""
    def fn(cur, run_id):
        n = 0
        for v in versiones:
            if v == version_nueva:
                continue
            cur.execute("UPDATE [int].dimPlanGeneracion SET EsVigente = 0, FechaCargaInt = SYSDATETIME(), RunId = ? "
                        "WHERE PlanVersion = ? AND EsVigente = 1", run_id, v)
            n += cur.rowcount
        return n, 0, n, 0
    return fn


def ejecutar_sp(sp):
    def fn(cur, run_id):
        cur.execute(f"EXEC {sp} @RunId = ?", run_id)
        return tuple(cur.fetchone())
    return fn


def main() -> int:
    parser = argparse.ArgumentParser(description="Carga masiva del plan de generacion de energia (Sheet1 de Calculos) desde Excel.")
    parser.add_argument("--archivo", required=True, help="Ruta del Excel.")
    parser.add_argument("--anio", required=True, type=int, help="Anio del plan (el archivo trae solo el mes).")
    parser.add_argument("--version", required=True, help="Identificador de la version del plan (ej. V1, 2026-09-24).")
    parser.add_argument("--reemplaza-version", nargs="+", default=[], metavar="VERSION",
                        help="Versiones anteriores que esta carga reemplaza por completo (quedan como historicas).")
    parser.add_argument("--limpiar", action="store_true",
                        help="Vacia dimPlanGeneracion (stg, int y dw, con su historial) antes de cargar. Destructivo.")
    parser.add_argument("--hoja", default=HOJA_DEFECTO, help=f"Hoja a leer (default: {HOJA_DEFECTO}).")
    parser.add_argument("--fuente", default=FUENTE_DEFECTO, help=f"Origen del plan (default: '{FUENTE_DEFECTO}').")
    parser.add_argument("--dry-run", action="store_true", help="Solo valida y muestra el resumen; no escribe en la base.")
    parser.add_argument("--env-file", default=str(ROOT / ".env"))
    args = parser.parse_args()

    ruta = Path(args.archivo)
    if not ruta.is_file():
        print(f"ERROR: no existe el archivo {ruta}", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(ruta, data_only=True)
    errores, avisos = [], []
    filas = leer_filas(wb, args.hoja, errores, avisos)
    plan = []
    if filas and not errores:
        evaluar_calidad(filas, args.anio, avisos)
        plan = expandir_dias(filas, args.anio)

    print(f"Archivo: {ruta.name} | hoja: {args.hoja} | anio: {args.anio} | version: {args.version} | fuente: {args.fuente}")
    if plan:
        inv = {(p["Proveedor"], p["InversorId"]) for p in plan}
        print(f"Filas del Excel: {len(filas)} | expandidas a dias: {len(plan)} | fechas: {plan[0]['Fecha']} a {plan[-1]['Fecha']} "
              f"| plantas: {len({p['Planta'] for p in plan})} | inversores: {len(inv)}")
        print(f"Plan de inversores: {sum(p['PlanDiarioInversor'] or 0 for p in plan):,.0f} kWh en el anio")

    conn = None
    if not errores and plan:
        conn = build_connection(load_env(Path(args.env_file)))
        conn.autocommit = False
        avisar_dispositivos(conn.cursor(), plan, avisos)

    for a in avisos:
        print(f"AVISO: {a}")
    if errores:
        for e in errores:
            print(f"ERROR: {e}", file=sys.stderr)
        print("No se escribio nada en la base.", file=sys.stderr)
        return 2
    if not plan:
        print("ERROR: no quedaron filas de plan para cargar.", file=sys.stderr)
        return 2
    if args.dry_run:
        conn.close()
        print("Dry-run: validacion correcta, no se escribio nada.")
        return 0

    try:
        print("Cargando:")
        nombre = ruta.name
        pasos = [
            ("PlanGeneracion_Bronze", ("Excel", args.hoja), ("stg", "dimPlanGeneracion"), cargar_stg(plan, nombre, args.fuente, args.version)),
            ("PlanGeneracion_Silver", ("stg", "dimPlanGeneracion"), ("int", "dimPlanGeneracion"), ejecutar_sp("[int].usp_MergeDimPlanGeneracion")),
            ("PlanGeneracion_Gold", ("int", "dimPlanGeneracion"), ("dw", "dimPlanGeneracion"), ejecutar_sp("dw.usp_MergeDimPlanGeneracion")),
        ]
        if args.limpiar:
            pasos.insert(0, ("PlanGeneracion_Limpieza", ("dw", "dimPlanGeneracion"), ("dw", "dimPlanGeneracion"), limpiar_tablas))
        if args.reemplaza_version:
            pasos.insert(len(pasos) - 1, ("PlanGeneracion_Reemplazo", ("int", "dimPlanGeneracion"), ("int", "dimPlanGeneracion"),
                                          retirar_versiones(args.reemplaza_version, args.version)))
        for proceso, origen, destino, fn in pasos:
            ejecutar_paso(conn, proceso, origen, destino, fn)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        conn.close()

    print("Carga del plan de generacion completada.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
