"""
Carga el catalogo de plantas del dominio Solar desde el Excel "Catalogo Plantas Solares" hacia
stg -> [int] -> dw.dimPlantaSolar / dw.dimPlantaSolarOrigen (migraciones 258-260).

Es lo unico manual del dominio (junto con el plan de generacion): los inversores, su numero de serie y
su capacidad salen solos de las dimensiones de cada portal (dw.vwSolarEquipo). El Excel solo dice lo
que ningun portal sabe:

  Hoja Plantas : Planta | Plantel | Orden | MeteoSite | EnReporte (SI/NO)
  Hoja Origen  : Proveedor | CodigoOrigen | NombreOrigen | Planta
                 Proveedor = sma, huawei, soliscloud o growatt. CodigoOrigen = agrupador del portal
                 (SMA PlantId, Huawei StationCode, Soliscloud StationId; en Growatt el serial del equipo).
                 NombreOrigen es solo referencia. Planta vacia = agrupador excluido a proposito.
  Otras hojas (ej. Equipos) se ignoran.

Valida antes de escribir; si hay errores no se escribe nada:
  * cabeceras; Planta unica y obligatoria; Orden entero; EnReporte SI/NO;
  * Proveedor valido, CodigoOrigen obligatorio y sin repetir por proveedor;
  * la Planta de Origen debe existir en la hoja Plantas.
Avisa, sin bloquear: MeteoSite sin datos en dw.factMeteoDaily, agrupadores que no existen en los portales,
agrupadores de los portales que no estan en Origen, plantas del informe sin inversores, y equipos del plan
vigente que no cruzarian. Muestra como quedaria cada planta (inversores y capacidad).

Carga FULL: lo que no viene en el Excel queda con EsVigente = 0 (no se borra).

Requiere openpyxl.

Uso:
    python db/etl/dimPlantaSolar/cargar_plantas_solares.py --archivo "<ruta>.xlsx" --dry-run
    python db/etl/dimPlantaSolar/cargar_plantas_solares.py --archivo "<ruta>.xlsx" [--env-file .env.prod]
"""
import argparse
import datetime
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent.parent.parent
sys.path.insert(0, str(ROOT / "db"))
from migrate import build_connection, load_env  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HOJA_PLANTAS = "Plantas"
HOJA_ORIGEN = "Origen"
COLS_PLANTAS = ["PLANTA", "PLANTEL", "ORDEN", "METEOSITE", "ENREPORTE"]
COLS_ORIGEN = ["PROVEEDOR", "CODIGOORIGEN", "NOMBREORIGEN", "PLANTA"]
PROVEEDORES = ("sma", "huawei", "soliscloud", "growatt")
SI = {"SI", "S", "1", "TRUE", "VERDADERO", "X"}
NO = {"NO", "N", "0", "FALSE", "FALSO", ""}


def norm(texto) -> str:
    s = unicodedata.normalize("NFD", str(texto or ""))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip().upper()


def texto(v) -> str:
    """Celda a texto; los numeros enteros de Excel (ej. 4551207.0) sin decimales."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def leer_hoja(wb, hoja, requeridas, errores) -> list:
    if hoja not in wb.sheetnames:
        errores.append(f"No existe la hoja '{hoja}'. Hojas: {wb.sheetnames}")
        return []
    filas = list(wb[hoja].iter_rows(values_only=True))
    if not filas:
        errores.append(f"La hoja '{hoja}' esta vacia.")
        return []
    cab = {norm(c): j for j, c in enumerate(filas[0]) if c is not None}
    faltan = [c for c in requeridas if c not in cab]
    if faltan:
        errores.append(f"Faltan las columnas {faltan} en la primera fila de '{hoja}'.")
        return []
    salida = []
    for n, r in enumerate(filas[1:], start=2):
        if r is None or all(v is None or str(v).strip() == "" for v in r):
            continue
        salida.append((n, {k: (r[cab[k]] if cab[k] < len(r) else None) for k in requeridas}))
    return salida


def validar_plantas(filas, errores) -> list:
    plantas, vistos, ordenes = [], {}, {}
    for n, r in filas:
        planta = texto(r["PLANTA"])
        if not planta:
            errores.append(f"'{HOJA_PLANTAS}' fila {n}: falta Planta.")
            continue
        if len(planta) > 30:
            errores.append(f"'{HOJA_PLANTAS}' fila {n}: Planta '{planta}' tiene mas de 30 caracteres.")
            continue
        if norm(planta) in vistos:
            errores.append(f"'{HOJA_PLANTAS}' fila {n}: la planta '{planta}' ya esta en la fila {vistos[norm(planta)]}.")
            continue
        vistos[norm(planta)] = n
        orden = r["ORDEN"]
        if orden in (None, ""):
            orden = None
        elif isinstance(orden, (int, float)) and not isinstance(orden, bool) and float(orden).is_integer():
            orden = int(orden)
        else:
            errores.append(f"'{HOJA_PLANTAS}' fila {n} ({planta}): Orden invalido {orden!r}.")
            continue
        en = norm(texto(r["ENREPORTE"]))
        if en not in SI | NO:
            errores.append(f"'{HOJA_PLANTAS}' fila {n} ({planta}): EnReporte debe ser SI o NO (viene {r['ENREPORTE']!r}).")
            continue
        en_reporte = en in SI
        meteo = texto(r["METEOSITE"])
        if en_reporte and not meteo:
            errores.append(f"'{HOJA_PLANTAS}' fila {n} ({planta}): falta MeteoSite (la planta va en el informe).")
            continue
        if en_reporte and orden is not None:
            ordenes.setdefault(orden, []).append(planta)
        plantas.append({"Planta": planta, "Plantel": texto(r["PLANTEL"]) or None, "Orden": orden,
                        "MeteoSite": meteo or None, "EnReporte": en_reporte})
    for orden, nombres in sorted(ordenes.items()):
        if len(nombres) > 1:
            errores.append(f"'{HOJA_PLANTAS}': el Orden {orden} se repite en {nombres}.")
    return plantas


def validar_origen(filas, plantas, errores) -> list:
    por_norm = {norm(p["Planta"]): p["Planta"] for p in plantas}
    origen, vistos = [], {}
    for n, r in filas:
        prov = texto(r["PROVEEDOR"]).lower()
        codigo = texto(r["CODIGOORIGEN"])
        if prov not in PROVEEDORES:
            errores.append(f"'{HOJA_ORIGEN}' fila {n}: Proveedor '{prov}' invalido (use {', '.join(PROVEEDORES)}).")
            continue
        if not codigo:
            errores.append(f"'{HOJA_ORIGEN}' fila {n}: falta CodigoOrigen.")
            continue
        if (prov, codigo) in vistos:
            errores.append(f"'{HOJA_ORIGEN}' fila {n}: {prov} {codigo} ya esta en la fila {vistos[(prov, codigo)]}.")
            continue
        vistos[(prov, codigo)] = n
        planta = texto(r["PLANTA"])
        if planta and norm(planta) not in por_norm:
            errores.append(f"'{HOJA_ORIGEN}' fila {n} ({prov} {codigo}): la planta '{planta}' no esta en la hoja {HOJA_PLANTAS}.")
            continue
        origen.append({"Proveedor": prov, "CodigoOrigen": codigo, "NombreOrigen": texto(r["NOMBREORIGEN"]) or None,
                       "Planta": por_norm[norm(planta)] if planta else None})
    return origen


def revisar_contra_base(cur, plantas, origen, avisos) -> None:
    """Avisos (no bloquean) y resumen de como quedaria cada planta."""
    cur.execute("SELECT DISTINCT Site FROM dw.factMeteoDaily")
    sitios = {str(r[0]).strip() for r in cur.fetchall()}
    for p in plantas:
        if p["MeteoSite"] and p["MeteoSite"] not in sitios:
            avisos.append(f"{p['Planta']}: el MeteoSite '{p['MeteoSite']}' no tiene datos en dw.factMeteoDaily (sin HSF ni PR).")

    cur.execute("SELECT Proveedor, CodigoOrigen, NombreOrigen, Serie, DeviceKey, NombrePortal, CapacidadDcPortalKwp FROM dw.vwSolarEquipo")
    equipos = [tuple(r) for r in cur.fetchall()]
    # misma regla que dw.vwRptInversor: portal; si no, dimDeviceCapacity por JoinKey (Huawei tambien por SourceSn = DeviceName)
    cur.execute("SELECT Vendor, JoinKey, SourceSn, CapacityDcKwp FROM dw.dimDeviceCapacity WHERE EsVigente = 1")
    externa = [tuple(r) for r in cur.fetchall()]

    def capacidad_externa(prov, key, nombre_portal):
        por_nombre = None
        for v, jk, sn, c in externa:
            if v != prov:
                continue
            if jk is not None and str(jk) == str(key):
                return c
            if prov == "huawei" and sn is not None and sn == nombre_portal and por_nombre is None:
                por_nombre = c
        return por_nombre

    grupos = {}
    for prov, cod, nombre, serie, key, nombre_portal, cap in equipos:
        cap = cap if cap is not None else capacidad_externa(prov, key, nombre_portal)
        grupos.setdefault((prov, cod), {"nombre": nombre, "equipos": []})["equipos"].append((serie, key, cap))

    asignado = {(o["Proveedor"], o["CodigoOrigen"]): o for o in origen}
    for clave in sorted(set(asignado) - set(grupos)):
        avisos.append(f"Origen {clave[0]} {clave[1]}: no existe en los portales (no aporta inversores).")
    for clave in sorted(set(grupos) - set(asignado)):
        g = grupos[clave]
        avisos.append(f"{clave[0]} {clave[1]} ({g['nombre'] or 'sin nombre'}, {len(g['equipos'])} inversor(es)) "
                      "no esta en la hoja Origen: no entra al informe y el monitor avisara.")

    print("\nComo quedaria el informe:")
    claves_plan = set()
    for p in sorted(plantas, key=lambda x: (x["Orden"] is None, x["Orden"] or 0, x["Planta"])):
        invs = []
        for (prov, cod), o in asignado.items():
            if o["Planta"] == p["Planta"] and (prov, cod) in grupos:
                for serie, key, cap in grupos[(prov, cod)]["equipos"]:
                    invs.append((prov, serie, key, cap))
                    claves_plan.add((prov, str(key)))
        cap_total = sum(c for *_, c in invs if c is not None)
        sin_cap = [s for _, s, _, c in invs if c is None]
        estado = "en informe" if p["EnReporte"] else "fuera del informe"
        print(f"  {str(p['Orden'] or '-'):>3} {p['Planta']:<14} {p['Plantel'] or '':<10} meteo={p['MeteoSite'] or '-':<14} "
              f"{len(invs):>2} inversores {cap_total:>9,.2f} kWp  ({estado})")
        if p["EnReporte"] and not invs:
            avisos.append(f"{p['Planta']}: va en el informe pero no tiene inversores (revisar la hoja Origen).")
        if p["EnReporte"] and sin_cap:
            avisos.append(f"{p['Planta']}: {len(sin_cap)} inversor(es) sin capacidad DC {sin_cap}.")

    cur.execute("SELECT DISTINCT Planta, InversorId, Proveedor, Inversor FROM dw.dimPlanGeneracion WHERE EsVigente = 1")
    sin_cruce = sorted(f"{r[0]} {r[1]}" for r in cur.fetchall() if (str(r[2]), str(r[3])) not in claves_plan)
    if sin_cruce:
        avisos.append(f"{len(sin_cruce)} inversor(es) del plan vigente no cruzarian con el catalogo y su plan no se reportaria: {sin_cruce}.")


# --- Carga (framework de control ETL) -----------------------------------------

def ejecutar_paso(conn, proceso, origen, destino, fn):
    cur = conn.cursor()
    cur.execute(
        """
        DECLARE @pid INT;
        EXEC dbo.usp_Etl_ProcesoRegistrar
            @Proceso = ?, @Dominio = ?, @SistemaOrigen = ?,
            @EsquemaOrigen = ?, @TablaOrigen = ?,
            @EsquemaDestino = ?, @TablaDestino = ?, @TipoCarga = ?,
            @ProcesoId = @pid OUTPUT;
        SELECT @pid;
        """,
        proceso, "Solar", origen[0], origen[0], origen[1], destino[0], destino[1], "FULL",
    )
    cur.fetchone()
    conn.commit()
    cur.execute("EXEC dbo.usp_Etl_RunIniciar @Proceso = ?", proceso)
    run_id = cur.fetchone()[0]
    conn.commit()
    try:
        leidas, insertadas, actualizadas, ignoradas = fn(cur, run_id)
        print(f"  [{proceso}] leidas {leidas} / insertadas {insertadas} / actualizadas {actualizadas} / sin cambio {ignoradas}")
        cur.execute(
            "EXEC dbo.usp_Etl_RunFinalizar @RunId = ?, @Estado = ?, @FilasLeidas = ?, @FilasInsertadas = ?, "
            "@FilasActualizadas = ?, @FilasIgnoradas = ?, @MensajeError = ?, @TareaError = ?",
            run_id, "EXITO", leidas, insertadas, actualizadas, ignoradas, None, None,
        )
        cur.execute(
            "EXEC dbo.usp_Etl_WatermarkActualizar @Proceso = ?, @NuevaFechaHora = ?, @TipoCarga = ?",
            proceso, datetime.datetime.now(), "FULL",
        )
        conn.commit()
    except Exception as exc:
        conn.rollback()
        cur.execute(
            "EXEC dbo.usp_Etl_RunFinalizar @RunId = ?, @Estado = ?, @FilasLeidas = ?, @FilasInsertadas = ?, "
            "@FilasActualizadas = ?, @FilasIgnoradas = ?, @MensajeError = ?, @TareaError = ?",
            run_id, "ERROR", None, None, None, None, str(exc), proceso,
        )
        conn.commit()
        raise


def cargar_stg(plantas, origen, archivo):
    def fn(cur, run_id):
        cur.execute("TRUNCATE TABLE stg.dimPlantaSolar")
        cur.execute("TRUNCATE TABLE stg.dimPlantaSolarOrigen")
        cur.executemany(
            "INSERT INTO stg.dimPlantaSolar (Planta, Plantel, Orden, MeteoSite, EnReporte, ArchivoOrigen, RunId) VALUES (?,?,?,?,?,?,?)",
            [(p["Planta"], p["Plantel"], p["Orden"], p["MeteoSite"], p["EnReporte"], archivo, run_id) for p in plantas],
        )
        cur.executemany(
            "INSERT INTO stg.dimPlantaSolarOrigen (Proveedor, CodigoOrigen, NombreOrigen, Planta, ArchivoOrigen, RunId) VALUES (?,?,?,?,?,?)",
            [(o["Proveedor"], o["CodigoOrigen"], o["NombreOrigen"], o["Planta"], archivo, run_id) for o in origen],
        )
        n = len(plantas) + len(origen)
        return n, n, 0, 0
    return fn


def ejecutar_sp(sp):
    def fn(cur, run_id):
        cur.execute(f"EXEC {sp} @RunId = ?", run_id)
        return tuple(cur.fetchone())
    return fn


def main() -> int:
    parser = argparse.ArgumentParser(description="Carga el catalogo de plantas solares (hojas Plantas y Origen) desde Excel.")
    parser.add_argument("--archivo", required=True, help="Ruta del Excel.")
    parser.add_argument("--dry-run", action="store_true", help="Solo valida y muestra como quedaria; no escribe en la base.")
    parser.add_argument("--env-file", default=str(ROOT / ".env"))
    args = parser.parse_args()

    ruta = Path(args.archivo)
    if not ruta.is_file():
        print(f"ERROR: no existe el archivo {ruta}", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(ruta, data_only=True)
    errores, avisos = [], []
    plantas = validar_plantas(leer_hoja(wb, HOJA_PLANTAS, COLS_PLANTAS, errores), errores)
    origen = validar_origen(leer_hoja(wb, HOJA_ORIGEN, COLS_ORIGEN, errores), plantas, errores)
    if not errores and not plantas:
        errores.append(f"La hoja '{HOJA_PLANTAS}' no tiene plantas.")

    print(f"Archivo: {ruta.name} | plantas: {len(plantas)} ({sum(p['EnReporte'] for p in plantas)} en el informe) "
          f"| agrupadores en Origen: {len(origen)} ({sum(1 for o in origen if o['Planta'] is None)} excluidos)")

    if errores:
        for e in errores:
            print(f"ERROR: {e}", file=sys.stderr)
        print("No se escribio nada en la base.", file=sys.stderr)
        return 2

    conn = build_connection(load_env(Path(args.env_file)))
    conn.autocommit = False
    try:
        revisar_contra_base(conn.cursor(), plantas, origen, avisos)
        for a in avisos:
            print(f"AVISO: {a}")
        if args.dry_run:
            print("Dry-run: validacion correcta, no se escribio nada.")
            return 0
        print("\nCargando:")
        pasos = [
            ("PlantaSolar_Bronze", ("Excel", "Plantas+Origen"), ("stg", "dimPlantaSolar"), cargar_stg(plantas, origen, ruta.name)),
            ("PlantaSolar_Silver", ("stg", "dimPlantaSolar"), ("int", "dimPlantaSolar"), ejecutar_sp("[int].usp_MergeDimPlantaSolar")),
            ("PlantaSolarOrigen_Silver", ("stg", "dimPlantaSolarOrigen"), ("int", "dimPlantaSolarOrigen"),
             ejecutar_sp("[int].usp_MergeDimPlantaSolarOrigen")),
            ("PlantaSolar_Gold", ("int", "dimPlantaSolar"), ("dw", "dimPlantaSolar"), ejecutar_sp("dw.usp_MergeDimPlantaSolar")),
            ("PlantaSolarOrigen_Gold", ("int", "dimPlantaSolarOrigen"), ("dw", "dimPlantaSolarOrigen"),
             ejecutar_sp("dw.usp_MergeDimPlantaSolarOrigen")),
        ]
        for proceso, desde, hacia, fn in pasos:
            ejecutar_paso(conn, proceso, desde, hacia, fn)
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1
    finally:
        conn.close()

    print("Carga del catalogo de plantas solares completada.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
